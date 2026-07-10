"""One-pager tests (offline): prompt building + JSON parsing + .docx rendering."""

import json
import zipfile
from datetime import date, datetime

import pytest
from docx import Document

from usage_agent.prompts.one_pager import (
    DEFAULT_ONE_PAGER_PROMPT,
    build_one_pager_prompt,
)
from usage_agent.reports.one_pager import (
    _delta,
    default_output_path,
    parse_report,
    render_one_pager_docx,
)

_SAMPLE_REPORT = {
    "title": "AI Usage & Cost Weekly Briefing",
    "period_label": "Jun 29 - Jul 5, 2026",
    "prior_period_label": "Jun 22 - Jun 28, 2026",
    "kpis": {
        "total_spend": {"last": 14200, "prior": 12650},
        "anthropic_spend": {"last": 9000, "prior": 8000},
        "openai_spend": {"last": 5200, "prior": 4650},
        "active_users": {"last": 120, "prior": 110},
    },
    "providers": [
        {"name": "Anthropic (Claude)", "last": 9000, "prior": 8000},
        {"name": "OpenAI (Codex)", "last": 5200, "prior": 4650},
    ],
    "models": [
        {"name": "claude-opus-4-8", "last": 6000, "prior": 5000},
        {"name": "gpt-5-codex", "last": 5200, "prior": 4650},
        {"name": "claude-sonnet-5", "last": 3000, "prior": 3000},
    ],
    "top_users": [
        {"name": "jane@x.com", "department": "Tax", "spend": 900, "top_product": "claude_code"},
        {"name": "bob@x.com", "department": "", "spend": 600, "top_product": "chatgpt"},
    ],
    "key_findings": ["Claude Code drove most of the growth."],
    "recommended_actions": ["Reclaim idle seats."],
}


# --- prompt building -------------------------------------------------------
def test_default_prompt_when_empty():
    assert build_one_pager_prompt("") == DEFAULT_ONE_PAGER_PROMPT
    assert build_one_pager_prompt(None) == DEFAULT_ONE_PAGER_PROMPT


def test_default_prompt_is_previous_week_wow():
    p = DEFAULT_ONE_PAGER_PROMPT
    assert "PREVIOUS COMPLETE WEEK" in p
    assert "PRIOR WEEK" in p
    assert "is_complete_week" in p
    assert "[Spend LW]" in p            # explicitly avoids the broken measures


def test_prompt_requires_json_shape():
    for p in (DEFAULT_ONE_PAGER_PROMPT, build_one_pager_prompt("claude spend last month")):
        assert "SINGLE JSON object" in p
        assert '"top_users"' in p       # the fixed schema is embedded


def test_user_prompt_kept_and_report_forced():
    out = build_one_pager_prompt("claude spend by department last month")
    assert out.startswith("claude spend by department last month")
    assert "SINGLE JSON object" in out


# --- JSON parsing ----------------------------------------------------------
def test_parse_report_plain_json():
    assert parse_report(json.dumps(_SAMPLE_REPORT))["title"] == _SAMPLE_REPORT["title"]


def test_parse_report_strips_fence_and_preamble():
    wrapped = 'Here is the briefing:\n```json\n' + json.dumps(_SAMPLE_REPORT) + "\n```"
    assert parse_report(wrapped)["period_label"] == "Jun 29 - Jul 5, 2026"


def test_parse_report_raises_without_json():
    with pytest.raises(ValueError):
        parse_report("no json here at all")


# --- WoW delta formatting --------------------------------------------------
def test_delta_rising_cost_is_attention_colored():
    # Spend up => "attention" (red), with signed $ and % (12.5 rounds to 12).
    text, color = _delta(9000, 8000, higher_is_better=False, money=True)
    assert color == "B23B3B" and text == "+$1,000 (+12%)"


def test_delta_more_users_is_positive_colored():
    text, color = _delta(120, 110, higher_is_better=True, money=False)
    assert color == "0F7A34" and text == "+10 (+9%)"


def test_delta_falling_cost_is_positive_colored():
    text, color = _delta(8000, 9000, higher_is_better=False, money=True)
    assert color == "0F7A34" and text.startswith("-$1,000")


def test_delta_new_when_prior_zero():
    text, _ = _delta(500, 0, higher_is_better=False, money=True)
    assert "(new)" in text


# --- output path -----------------------------------------------------------
def test_default_output_path_has_date():
    p = default_output_path(date(2026, 7, 9))
    assert p.name == "ai_usage_one_pager_2026-07-09.docx"


# --- docx rendering --------------------------------------------------------
def test_render_produces_valid_docx(tmp_path):
    out = tmp_path / "one_pager.docx"
    written = render_one_pager_docx(
        _SAMPLE_REPORT, out, generated_at=datetime(2026, 7, 6, 9, 30)
    )
    assert written == out.resolve()
    assert zipfile.is_zipfile(out)        # .docx is a zip package
    doc = Document(str(out))
    body = "\n".join(p.text for p in doc.paragraphs)
    cells = "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    footer = "\n".join(p.text for p in doc.sections[0].footer.paragraphs)
    assert "AI Usage & Cost Weekly Briefing" in body
    # The generated timestamp shows in both the masthead (a table) and the footer.
    assert "Jul 06, 2026 09:30" in (cells + footer)


def test_render_includes_all_three_tables(tmp_path):
    out = tmp_path / "op.docx"
    render_one_pager_docx(_SAMPLE_REPORT, out)
    doc = Document(str(out))
    headings = [p.text for p in doc.paragraphs] + [
        c.text for t in doc.tables for row in t.rows for c in row.cells
    ]
    assert any("Spend by provider" in h for h in headings)
    assert any("Spend by model" in h for h in headings)
    assert any("Top spenders" in h for h in headings)
    # masthead + KPI tiles + provider + model + top-spenders + findings/actions.
    assert len(doc.tables) == 6

    cells = [c.text for t in doc.tables for row in t.rows for c in row.cells]
    assert "jane@x.com" in cells
    assert "+$1,000 (+12%)" in cells       # provider table still shows WoW
    # Top-spenders table shows spend + most-used product, and NO WoW for users.
    assert "claude_code" in cells
    assert not any("$900" in c and "(" in c for c in cells)   # no "(+x%)" beside a user's spend


def test_render_creates_missing_parent_dirs(tmp_path):
    out = tmp_path / "nested" / "dir" / "op.docx"
    render_one_pager_docx(_SAMPLE_REPORT, out)
    assert out.exists()


def test_render_tolerates_sparse_report(tmp_path):
    # Missing sections must not crash — the layout stays fixed with empty tables.
    out = tmp_path / "sparse.docx"
    render_one_pager_docx({"title": "AI Usage & Cost Weekly Briefing"}, out)
    assert zipfile.is_zipfile(out)
